# build_docx.py —— 从 LaTeX 源（by1-en.tex / by1-zh.tex + by1-refs.bib）
# 生成 Word 版：paper/by1-en.docx、paper/by1-zh.docx。
# 用法：<捆绑 python> paper/build_docx.py
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent


# ── 小工具 ────────────────────────────────────────────────────────────
def read_brace(s, i):
    """s[i] == '{'；返回 (内容, 下一索引)。"""
    depth = 0
    j = i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


def eq_arg(s, name):
    """找 name = {...}（lstlisting 选项）并返回内容。"""
    m = re.search(re.escape(name) + r"\s*=\s*", s)
    if not m:
        return None
    j = m.end()
    if j < len(s) and s[j] == "{":
        arg, _ = read_brace(s, j)
        return arg
    return ""


def macro_arg(s, name, start=0):
    """找反斜杠 name{...} 并返回 (内容, 结束索引)。"""
    key = "\\" + name
    k = s.find(key, start)
    if k < 0:
        return None, start
    j = k + len(key)
    while j < len(s) and s[j] in " \t\n":
        j += 1
    if j < len(s) and s[j] == "{":
        arg, j2 = read_brace(s, j)
        return arg, j2
    return "", j


# ── bib 解析 ──────────────────────────────────────────────────────────
def parse_bib(path):
    s = path.read_text(encoding="utf-8")
    entries = {}
    for m in re.finditer(r"@(\w+)\s*\{", s):
        etype = m.group(1).lower()
        if etype == "comment":
            continue
        j = m.end() - 1
        key_body, k = read_brace(s, j)
        km = re.match(r"\s*([^,]+),", key_body)
        if not km:
            continue
        key = km.group(1).strip()
        body = key_body[km.end():]
        fields = {}
        i = 0
        while i < len(body):
            fm = re.compile(r"([a-zA-Z]+)\s*=").search(body, i)
            if not fm:
                break
            name = fm.group(1).lower()
            j2 = fm.end()
            while j2 < len(body) and body[j2] in " \t\n":
                j2 += 1
            if j2 < len(body) and body[j2] == "{":
                val, j3 = read_brace(body, j2)
            elif j2 < len(body) and body[j2] == '"':
                j3 = body.find('"', j2 + 1)
                val = body[j2 + 1:j3] if j3 > 0 else body[j2 + 1:]
                j3 = j3 + 1 if j3 > 0 else len(body)
            else:
                i = j2
                continue
            fields[name] = val.strip()
            i = j3
        entries[key] = {"type": etype, "fields": fields}
    return entries


def clean_tex(s):
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\textit\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\emph\{([^}]*)\}", r"\1", s)
    s = s.replace("{", "").replace("}", "")
    s = re.sub(r"\\['\^\"\~]\s*([a-zA-Z])", r"\1", s)
    s = s.replace("\\&", "&")
    return re.sub(r"\s+", " ", s).strip()


def author_list(raw):
    raw = raw.strip()
    while raw.startswith("{") and raw.endswith("}"):
        raw = raw[1:-1]
    parts = [clean_tex(p) for p in re.split(r"\s+and\s+", raw)]
    return [p for p in parts if p]


def format_ref(key, e):
    f = e["fields"]
    authors = author_list(f.get("author", ""))
    if len(authors) > 1:
        aup = ", ".join(authors[:-1]) + ", & " + authors[-1]
    else:
        aup = authors[0] if authors else ""
    year = f.get("year", "")
    title = clean_tex(f.get("title", ""))
    venue = clean_tex(f.get("journal") or f.get("booktitle") or "")
    vol = clean_tex(f.get("volume", ""))
    if vol:
        venue = (venue + ", " + vol).strip(", ")
    note = clean_tex(f.get("note", ""))
    how = clean_tex(f.get("howpublished", ""))
    tail = ". ".join(x for x in (venue, note, how) if x)
    out = f"{aup} ({year}). {title}."
    if tail:
        out += " " + tail + "."
    return out


# ── 行内渲染 ──────────────────────────────────────────────────────────
def parse_inline(text, ctx, fmt=()):
    segs = []
    fmt = tuple(fmt)

    def emit(t, f):
        if t:
            segs.append((re.sub(r"[ \t]+", " ", t), f))

    i = 0
    while i < len(text):
        c = text[i]
        if text.startswith("---", i):
            emit("—", fmt); i += 3; continue
        if text.startswith("--", i):
            emit("–", fmt); i += 2; continue
        if c == "\\" and i + 1 < len(text):
            n = text[i + 1]
            if n in "%_&#$":
                emit(n, fmt); i += 2; continue
            if n == "," or n == " ":
                emit(" ", fmt); i += 2; continue
            m = re.compile(
                r"\\(code|texttt|textbf|emph|q|byone|ir|cite|ref|label|url|textsc)"
            ).match(text, i)
            if m:
                name = m.group(1)
                k = m.end()
                if k < len(text) and text[k] == "{":
                    arg, k2 = read_brace(text, k)
                else:
                    arg, k2 = "", k
                if name in ("code", "texttt"):
                    segs += parse_inline(arg, ctx, fmt + ("mono",))
                elif name == "textbf":
                    segs += parse_inline(arg, ctx, fmt + ("b",))
                elif name == "emph":
                    segs += parse_inline(arg, ctx, fmt + ("i",))
                elif name == "q":
                    lq, rq = ("「", "」") if ctx["lang"] == "zh" else ("“", "”")
                    emit(lq, fmt); segs += parse_inline(arg, ctx, fmt); emit(rq, fmt)
                elif name == "byone":
                    emit("by1", fmt)
                elif name == "ir":
                    emit("by1-ir", fmt)
                elif name == "cite":
                    emit(cite_text(arg, ctx), fmt)
                elif name == "ref":
                    emit(ctx["labels"].get(arg.strip(), "?"), fmt)
                elif name in ("label",):
                    pass
                else:
                    segs += parse_inline(arg, ctx, fmt)
                i = k2
                continue
            emit("\\", fmt); i += 1; continue
        if c == "$":
            j = text.find("$", i + 1)
            if j > 0:
                math = text[i + 1:j]
                math = math.replace("\\times", "×").replace("\\,", " ").replace("\\;", " ")
                emit(math, fmt)
                i = j + 1
                continue
        if c == "~":
            emit(" ", fmt); i += 1; continue
        if c in "{}":
            i += 1
            continue
        if c == "\n":
            emit(" ", fmt); i += 1
            continue
        emit(c, fmt); i += 1
    return segs


def cite_text(arg, ctx):
    keys = [k.strip() for k in arg.split(",") if k.strip()]
    nums = [str(ctx["cite_no"].get(k, "?")) for k in keys]
    return "[" + ",".join(nums) + "]"


# ── 结构解析 ──────────────────────────────────────────────────────────
def parse_tex(path):
    s = path.read_text(encoding="utf-8")
    title, _ = macro_arg(s, "title")
    author, _ = macro_arg(s, "author")
    date, _ = macro_arg(s, "date")
    body = s.split("\\begin{document}", 1)[1].split("\\end{document}")[0]
    lines = body.split("\n")

    blocks = []
    i = 0
    while i < len(lines):
        st = lines[i].strip()
        if not st or st == "\\maketitle":
            i += 1
            continue
        if st.startswith("\\begin{lstlisting}"):
            opt = lines[i]
            k = i
            while not opt.rstrip().endswith("]") and k + 1 < len(lines):
                k += 1
                opt += " " + lines[k].strip()
            cap = eq_arg(opt, "caption")
            lab = eq_arg(opt, "label")
            code = []
            i = k + 1
            while i < len(lines) and lines[i].strip() != "\\end{lstlisting}":
                code.append(lines[i])
                i += 1
            blocks.append({"type": "listing", "caption": cap, "label": lab, "code": code})
            i += 1
            continue
        if st.startswith("\\begin{abstract}"):
            buf = []
            i += 1
            while i < len(lines) and lines[i].strip() != "\\end{abstract}":
                buf.append(lines[i])
                i += 1
            blocks.append({"type": "abstract", "text": "\n".join(buf)})
            i += 1
            continue
        if st.startswith("\\begin{table}"):
            buf = []
            i += 1
            while i < len(lines) and lines[i].strip() != "\\end{table}":
                buf.append(lines[i])
                i += 1
            blocks.append({"type": "table", "text": "\n".join(buf)})
            i += 1
            continue
        if st.startswith("\\begin{itemize}") or st.startswith("\\begin{enumerate}"):
            ordered = st.startswith("\\begin{enumerate}")
            endtag = "\\end{enumerate}" if ordered else "\\end{itemize}"
            items = []
            cur = None
            i += 1
            while i < len(lines) and lines[i].strip() != endtag:
                t = lines[i].strip()
                if t.startswith("\\item"):
                    if cur is not None:
                        items.append(cur)
                    cur = t[5:].strip()
                elif cur is not None:
                    cur += " " + t
                i += 1
            if cur is not None:
                items.append(cur)
            blocks.append({"type": "list", "ordered": ordered, "items": items})
            i += 1
            continue
        if st.startswith("\\section{"):
            arg, _ = macro_arg(st, "section")
            blocks.append({"type": "section", "title": arg, "label": ""})
            i += 1
            continue
        if st.startswith("\\subsection{"):
            arg, _ = macro_arg(st, "subsection")
            blocks.append({"type": "subsection", "title": arg, "label": ""})
            i += 1
            continue
        if st.startswith("\\paragraph{"):
            arg, j2 = macro_arg(st, "paragraph")
            buf = []
            rest = st[j2:].strip()
            if rest:
                buf.append(rest)
            i += 1
            while i < len(lines) and lines[i].strip():
                buf.append(lines[i].strip())
                i += 1
            blocks.append({"type": "para_head", "title": arg, "text": " ".join(buf)})
            continue
        if st.startswith("\\label{"):
            lab, _ = macro_arg(st, "label")
            if blocks:
                blocks[-1]["label"] = lab
            i += 1
            continue
        if st.startswith("\\bibliographystyle") or st.startswith("\\bibliography{"):
            i += 1
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"\\(begin|end|section|subsection|paragraph|item|label|bibliography|maketitle)",
            lines[i].strip(),
        ):
            buf.append(lines[i].strip())
            i += 1
        if buf:
            blocks.append({"type": "para", "text": " ".join(buf)})
    return {"title": title, "author": author, "date": date, "blocks": blocks}


def parse_table_block(text):
    cap, _ = macro_arg(text, "caption")
    lab, _ = macro_arg(text, "label")
    k = text.find("\\begin{tabular}")
    spec = ""
    rows = []
    if k >= 0:
        j = text.find("{", k + len("\\begin{tabular}"))
        spec, k2 = read_brace(text, j)
        end = text.find("\\end{tabular}", k2)
        content = text[k2:end]
        content = re.sub(r"\\(toprule|midrule|bottomrule)", "\n", content)
        for row in re.split(r"\\\\", content):
            row = row.strip()
            if not row:
                continue
            cells = [c.strip() for c in row.split("&")]
            rows.append(cells)
    return cap, lab, spec, rows


def col_widths(spec, n, total_cm=15.2):
    toks = re.findall(r"p\{([0-9.]+)cm\}|([lrc])", spec)
    fixed = [float(a) if a else None for a, b in toks]
    if len(fixed) == n and all(w for w in fixed):
        return fixed
    if len(fixed) == n:
        rest = total_cm - sum(w for w in fixed if w)
        cnt = sum(1 for w in fixed if not w)
        return [w if w else rest / cnt for w in fixed]
    return [total_cm / n] * n


# ── Word 渲染 ─────────────────────────────────────────────────────────
def style_run(run, fmt=(), size=None):
    mono = "mono" in fmt
    run.font.name = "Consolas" if mono else "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "SimSun")
    if mono:
        run.font.size = Pt(size or 9)
    if "b" in fmt:
        run.bold = True
    if "i" in fmt:
        run.italic = True
    if size and not mono:
        run.font.size = Pt(size)


def add_runs(par, segs, size=None):
    for t, f in segs:
        r = par.add_run(t)
        style_run(r, f, size)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11.5)):
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimHei")
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)


def three_line_table(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, sz in (("top", "12"), ("bottom", "12")):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        bd = OxmlElement("w:tcBorders")
        b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "6")
        b.set(qn("w:color"), "000000")
        bd.append(b)
        tcPr.append(bd)


def shade_paragraph(par, fill="F2F2F2"):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def tab_label(lang, num):
    return ("表 " + str(num) + "： ") if lang == "zh" else ("Table " + str(num) + ": ")


def lst_label(lang, num):
    return ("代码 " + str(num) + "： ") if lang == "zh" else ("Listing " + str(num) + ": ")


def build_docx(tex_path, bib, out_path, lang):
    data = parse_tex(tex_path)
    ctx = {"lang": lang, "labels": {}, "bib": bib}
    cite_no = {}
    for n, k in enumerate(sorted(bib.keys()), 1):
        cite_no[k] = n
    ctx["cite_no"] = cite_no

    sec_no, sub_no, tab_no, lst_no = 0, 0, 0, 0
    for b in data["blocks"]:
        if b["type"] == "section":
            sec_no += 1
            sub_no = 0
            b["num"] = str(sec_no)
            if b.get("label"):
                ctx["labels"][b["label"]] = b["num"]
        elif b["type"] == "subsection":
            sub_no += 1
            b["num"] = "%d.%d" % (sec_no, sub_no)
            if b.get("label"):
                ctx["labels"][b["label"]] = b["num"]
        elif b["type"] == "table":
            tab_no += 1
            cap, lab, spec, rows = parse_table_block(b["text"])
            b.update(cap=cap, spec=spec, rows=rows, num=tab_no)
            if lab:
                ctx["labels"][lab] = str(tab_no)
        elif b["type"] == "listing":
            lst_no += 1
            b["num"] = lst_no
            if b.get("label"):
                ctx["labels"][b["label"]] = str(lst_no)

    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Cm(2.6))

    for part in data["title"].split("\\\\"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, parse_inline(part.strip(), ctx, ("b",)), 16)
        p.paragraph_format.space_after = Pt(2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, parse_inline(data["author"], ctx), 11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, parse_inline(data["date"], ctx), 11)
    p.paragraph_format.space_after = Pt(12)

    abs_lbl = "摘要" if lang == "zh" else "Abstract"
    ref_lbl = "参考文献" if lang == "zh" else "References"

    for b in data["blocks"]:
        t = b["type"]
        if t == "abstract":
            p = doc.add_paragraph()
            r = p.add_run(abs_lbl)
            style_run(r, ("b",), 12)
            p.paragraph_format.space_before = Pt(8)
            for para_text in b["text"].split("\n\n"):
                joined = " ".join(x.strip() for x in para_text.split("\n") if x.strip())
                if not joined:
                    continue
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs(p, parse_inline(joined, ctx))
                for r in p.runs:
                    r.italic = True
        elif t == "section":
            p = doc.add_paragraph(style="Heading 1")
            add_runs(p, parse_inline(b["num"] + "  " + b["title"], ctx, ("b",)), 14)
        elif t == "subsection":
            p = doc.add_paragraph(style="Heading 2")
            add_runs(p, parse_inline(b["num"] + "  " + b["title"], ctx, ("b",)), 12)
        elif t == "para_head":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, parse_inline(b["title"] + " ", ctx, ("b",)))
            add_runs(p, parse_inline(b["text"], ctx))
            if lang == "zh":
                p.paragraph_format.first_line_indent = Cm(0.74)
        elif t == "para":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, parse_inline(b["text"], ctx))
            if lang == "zh":
                p.paragraph_format.first_line_indent = Cm(0.74)
        elif t == "list":
            style = "List Number" if b["ordered"] else "List Bullet"
            for it in b["items"]:
                p = doc.add_paragraph(style=style)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_runs(p, parse_inline(it, ctx))
        elif t == "table":
            p = doc.add_paragraph()
            r = p.add_run(tab_label(lang, b["num"]))
            style_run(r, ("b",), 10)
            add_runs(p, parse_inline(b["cap"], ctx), 10)
            p.paragraph_format.space_after = Pt(3)
            rows = b["rows"]
            if not rows:
                continue
            ncol = max(len(rw) for rw in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = col_widths(b["spec"], ncol)
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = table.cell(ri, ci)
                    cell.width = Cm(widths[ci] if ci < len(widths) else 3.0)
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    segs = parse_inline(row[ci] if ci < len(row) else "", ctx, ("b",) if ri == 0 else ())
                    add_runs(par, segs, 9.5)
            three_line_table(table)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "listing":
            p = doc.add_paragraph()
            r = p.add_run(lst_label(lang, b["num"]))
            style_run(r, ("b",), 10)
            add_runs(p, parse_inline(b["caption"], ctx), 10)
            for code_line in b["code"]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(code_line.rstrip() if code_line.strip() else " ")
                style_run(r, ("mono",), 9)
                shade_paragraph(p)

    p = doc.add_paragraph(style="Heading 1")
    add_runs(p, parse_inline(ref_lbl, ctx, ("b",)), 14)
    for n, k in enumerate(sorted(bib.keys()), 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        add_runs(p, parse_inline("[%d] " % n + format_ref(k, bib[k]), ctx), 10)

    doc.save(out_path)
    return out_path


def main():
    bib = parse_bib(HERE / "by1-refs.bib")
    print("bib entries:", len(bib))
    for tex, lang in (("by1-en.tex", "en"), ("by1-zh.tex", "zh")):
        out = HERE / (Path(tex).stem + ".docx")
        build_docx(HERE / tex, bib, out, lang)
        print("wrote", out, out.stat().st_size, "bytes")


if __name__ == "__main__":
    main()
